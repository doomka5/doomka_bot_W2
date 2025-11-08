from __future__ import annotations

import os
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlencode

import asyncpg
from fastapi import FastAPI, Form, HTTPException, Request, Response, status
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook

BASE_DIR = Path(__file__).resolve().parent

app = FastAPI(title="Doomka W2 — CRM склад")
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")

templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


def _format_value(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, Decimal):
        normalized = value.normalize()
        return format(normalized, "f").rstrip("0").rstrip(".") if "." in format(normalized, "f") else format(normalized, "f")
    return str(value)


templates.env.filters["format_value"] = _format_value


def url_with_query(request: Request, route_name: str, **params: Any) -> str:
    base_url = request.url_for(route_name)
    query = {key: value for key, value in params.items() if value not in (None, "")}
    if query:
        return f"{base_url}?{urlencode(query)}"
    return base_url


templates.env.globals["url_with_query"] = url_with_query

DB_SETTINGS = {
    "user": os.getenv("DB_USER", "botuser"),
    "password": os.getenv("DB_PASS", "botpass"),
    "database": os.getenv("DB_NAME", "botdb"),
    "host": os.getenv("DB_HOST", "postgres_bot"),
    "port": int(os.getenv("DB_PORT", 5432)),
}

TABLE_ALIASES: dict[str, list[str]] = {
    "materials": ["materials", "warehouse_plastics"],
    "films": ["films", "warehouse_films"],
}

ROUTE_NAMES: dict[str, str] = {
    "materials": "materials_page",
    "films": "films_page",
}

POWER_SUPPLY_COLUMNS: list[dict[str, Any]] = [
    {"name": "article", "title": "Артикул", "data_type": "text"},
    {"name": "manufacturer", "title": "Производитель", "data_type": "text"},
    {"name": "series", "title": "Серия", "data_type": "text"},
    {"name": "power", "title": "Мощность", "data_type": "text"},
    {"name": "voltage", "title": "Напряжение", "data_type": "text"},
    {"name": "ip", "title": "Степень защиты IP", "data_type": "text"},
    {"name": "total_quantity", "title": "Количество на складе", "data_type": "integer"},
]

POWER_SUPPLY_SUMMARY_CTE = """
WITH summary AS (
    SELECT
        gps.article,
        manufacturer.name AS manufacturer,
        series.name AS series,
        power.name AS power,
        voltage.name AS voltage,
        ip.name AS ip,
        COALESCE(SUM(wps.quantity), 0) AS total_quantity
    FROM generated_power_supplies AS gps
    JOIN power_supply_manufacturers AS manufacturer ON manufacturer.id = gps.manufacturer_id
    JOIN power_supply_series AS series ON series.id = gps.series_id
    JOIN power_supply_power_options AS power ON power.id = gps.power_option_id
    JOIN power_supply_voltage_options AS voltage ON voltage.id = gps.voltage_option_id
    JOIN power_supply_ip_options AS ip ON ip.id = gps.ip_option_id
    LEFT JOIN warehouse_power_supplies AS wps ON wps.power_supply_id = gps.id
    GROUP BY
        gps.id,
        gps.article,
        manufacturer.name,
        series.name,
        power.name,
        voltage.name,
        ip.name
    HAVING COALESCE(SUM(wps.quantity), 0) > 0
)
"""

TEXT_TYPES = {"character varying", "text", "citext"}
INT_TYPES = {"smallint", "integer", "bigint"}
FLOAT_TYPES = {"double precision", "real"}
DECIMAL_TYPES = {"numeric", "decimal", "money"}
DATE_TYPES = {"date"}
TIME_TYPES = {"time without time zone", "time with time zone"}
TIMESTAMP_TYPES = {"timestamp without time zone", "timestamp with time zone"}
BOOL_TYPES = {"boolean"}

COLUMN_LABELS = {
    "id": "ID",
    "article": "Артикул",
    "material": "Материал",
    "thickness": "Толщина",
    "color": "Цвет",
    "color_code": "Код цвета",
    "length": "Длина",
    "width": "Ширина",
    "warehouse": "Склад",
    "comment": "Комментарий",
    "employee_id": "ID сотрудника",
    "employee_name": "Сотрудник",
    "employee_nick": "Ник сотрудника",
    "arrival_date": "Дата поступления",
    "arrival_at": "Дата и время поступления",
    "manufacturer": "Производитель",
    "series": "Серия",
    "recorded_at": "Зарегистрировано",
}

db_pool: asyncpg.Pool | None = None


def quote_ident(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


async def resolve_table(conn: asyncpg.Connection, alias: str) -> dict[str, str]:
    candidates = TABLE_ALIASES.get(alias)
    if not candidates:
        raise HTTPException(status_code=404, detail="Неизвестный раздел")

    for candidate in candidates:
        schema = "public"
        table = candidate
        if "." in candidate:
            schema, table = candidate.split(".", 1)
        exists = await conn.fetchval(
            """
            SELECT EXISTS(
                SELECT 1 FROM information_schema.tables
                WHERE table_schema = $1 AND table_name = $2
            )
            """,
            schema,
            table,
        )
        if exists:
            return {
                "schema": schema,
                "table": table,
                "qualified": f"{quote_ident(schema)}.{quote_ident(table)}",
                "alias": alias,
            }
    raise HTTPException(status_code=404, detail=f"Таблица для раздела '{alias}' не найдена")


async def fetch_columns(conn: asyncpg.Connection, schema: str, table: str) -> list[dict[str, Any]]:
    rows = await conn.fetch(
        """
        SELECT column_name, data_type, is_nullable, column_default
        FROM information_schema.columns
        WHERE table_schema = $1 AND table_name = $2
        ORDER BY ordinal_position
        """,
        schema,
        table,
    )
    if not rows:
        raise HTTPException(status_code=404, detail=f"Нет колонок у таблицы {schema}.{table}")

    columns: list[dict[str, Any]] = []
    for row in rows:
        name = row["column_name"]
        label = COLUMN_LABELS.get(name, name.replace("_", " ").capitalize())
        columns.append(
            {
                "name": name,
                "data_type": row["data_type"],
                "is_nullable": row["is_nullable"] == "YES",
                "has_default": row["column_default"] is not None,
                "title": label,
            }
        )
    return columns


async def fetch_rows(
    conn: asyncpg.Connection,
    table_info: dict[str, str],
    columns: list[dict[str, Any]],
    search: str | None,
    sort: str | None,
    order: str | None,
    filters: dict[str, str] | None = None,
) -> tuple[list[dict[str, Any]], str, str]:
    column_names = [col["name"] for col in columns]
    if not column_names:
        return [], "", "asc"

    sort_column = sort if sort in column_names else column_names[0]
    order_direction = "desc" if order and order.lower() == "desc" else "asc"

    text_columns = [col for col in columns if col["data_type"] in TEXT_TYPES]

    column_map = {col["name"]: col for col in columns}

    params: list[Any] = []
    conditions: list[str] = []
    param_index = 1

    if search and text_columns:
        like_parts = []
        for col in text_columns:
            like_parts.append(f"{quote_ident(col['name'])} ILIKE ${param_index}")
            params.append(f"%{search}%")
            param_index += 1
        if like_parts:
            conditions.append("(" + " OR ".join(like_parts) + ")")

    if filters:
        for column_name, raw_value in filters.items():
            if raw_value in (None, ""):
                continue
            column = column_map.get(column_name)
            if not column:
                continue
            try:
                typed_value = convert_value(raw_value, column["data_type"])
            except (TypeError, ValueError):
                continue
            conditions.append(f"{quote_ident(column_name)} = ${param_index}")
            params.append(typed_value)
            param_index += 1

    where_clause = ""
    if conditions:
        where_clause = " WHERE " + " AND ".join(conditions)

    select_columns = ", ".join(quote_ident(name) for name in column_names)
    query = (
        f"SELECT {select_columns} FROM {table_info['qualified']}"
        f"{where_clause} ORDER BY {quote_ident(sort_column)} {order_direction.upper()}"
    )

    records = await conn.fetch(query, *params)
    rows = [dict(record) for record in records]
    return rows, sort_column, order_direction


def convert_value(value: Any, data_type: str) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == "":
            return None
        value = stripped

    if data_type in INT_TYPES:
        return int(value)
    if data_type in DECIMAL_TYPES:
        return Decimal(str(value))
    if data_type in FLOAT_TYPES:
        return float(value)
    if data_type in DATE_TYPES:
        if isinstance(value, date) and not isinstance(value, datetime):
            return value
        return date.fromisoformat(str(value))
    if data_type in TIMESTAMP_TYPES:
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        return datetime.fromisoformat(str(value))
    if data_type in TIME_TYPES:
        if isinstance(value, time):
            return value
        return time.fromisoformat(str(value))
    if data_type in BOOL_TYPES:
        if isinstance(value, bool):
            return value
        lowered = str(value).strip().lower()
        if lowered in {"1", "true", "t", "yes", "y"}:
            return True
        if lowered in {"0", "false", "f", "no", "n"}:
            return False
        raise ValueError("ожидалось булево значение (true/false)")
    return str(value)


async def get_pool() -> asyncpg.Pool:
    if db_pool is None:
        raise HTTPException(status_code=500, detail="База данных недоступна")
    return db_pool


@app.on_event("startup")
async def startup() -> None:
    global db_pool
    db_pool = await asyncpg.create_pool(**DB_SETTINGS)


@app.on_event("shutdown")
async def shutdown() -> None:
    global db_pool
    if db_pool is not None:
        await db_pool.close()
        db_pool = None


@app.get("/", response_class=HTMLResponse)
async def index(request: Request) -> HTMLResponse:
    return templates.TemplateResponse("index.html", {"request": request})


def _prepare_form_columns(raw_columns: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    prepared: list[dict[str, Any]] = []
    for column in raw_columns:
        name = column["name"]
        if name == "id":
            continue
        data_type = column["data_type"]
        input_type = "text"
        step = None
        if data_type in INT_TYPES:
            input_type = "number"
            step = "1"
        elif data_type in DECIMAL_TYPES | FLOAT_TYPES:
            input_type = "number"
            step = "0.01"
        elif data_type in DATE_TYPES:
            input_type = "date"
        elif data_type in TIMESTAMP_TYPES:
            input_type = "datetime-local"
        elif data_type in TIME_TYPES:
            input_type = "time"
        elif data_type == "text":
            input_type = "textarea"
        prepared.append(
            {
                "name": name,
                "title": column["title"],
                "input_type": input_type,
                "step": step,
                "required": not column["is_nullable"] and not column["has_default"],
            }
        )
    return prepared


async def _render_table(
    request: Request,
    alias: str,
    template_name: str,
    search: str,
    sort: str | None,
    order: str | None,
    message: str | None,
    error: str | None,
    *,
    filters: dict[str, str] | None = None,
    filter_columns: list[str] | None = None,
) -> HTMLResponse:
    active_filters = dict(filters or {})
    filter_options: dict[str, list[dict[str, str]]] = {}

    pool = await get_pool()
    async with pool.acquire() as conn:
        table_info = await resolve_table(conn, alias)
        columns = await fetch_columns(conn, table_info["schema"], table_info["table"])
        column_names = [col["name"] for col in columns]

        if filter_columns:
            for column_name in filter_columns:
                if column_name not in column_names:
                    continue
                records = await conn.fetch(
                    f"SELECT DISTINCT {quote_ident(column_name)} FROM {table_info['qualified']}"
                    f" WHERE {quote_ident(column_name)} IS NOT NULL"
                    f" ORDER BY {quote_ident(column_name)} ASC"
                )
                options: list[dict[str, str]] = []
                for record in records:
                    value = record[column_name]
                    if value is None:
                        continue
                    options.append({
                        "value": str(value),
                        "label": _format_value(value),
                    })
                filter_options[column_name] = options

        rows, applied_sort, applied_order = await fetch_rows(
            conn,
            table_info,
            columns,
            search,
            sort,
            order,
            filters=active_filters,
        )

    export_url = request.url_for("export_table", table_alias=alias)
    query_params: dict[str, str] = {}
    if search:
        query_params["search"] = search
    if applied_sort:
        query_params["sort"] = applied_sort
    if applied_order:
        query_params["order"] = applied_order
    for key, value in active_filters.items():
        if value not in (None, ""):
            query_params[key] = value
    if query_params:
        export_url = f"{export_url}?{urlencode(query_params)}"

    route_name = ROUTE_NAMES.get(alias, f"{alias}_page")

    context = {
        "request": request,
        "columns": columns,
        "rows": rows,
        "search": search,
        "sort": applied_sort,
        "order": applied_order,
        "message": message,
        "error": error,
        "export_url": export_url,
        "route_name": route_name,
        "filters": active_filters,
        "filter_options": filter_options,
    }
    return templates.TemplateResponse(template_name, context)


async def fetch_power_supply_summary(
    conn: asyncpg.Connection,
    search: str,
    sort: str | None,
    order: str | None,
    filters: dict[str, str] | None = None,
) -> tuple[list[dict[str, Any]], str, str]:
    available_columns = {column["name"] for column in POWER_SUPPLY_COLUMNS}
    sort_column = sort if sort in available_columns else "total_quantity"

    if order is None:
        order_direction = "desc" if sort_column == "total_quantity" else "asc"
    else:
        order_direction = "desc" if order.lower() == "desc" else "asc"

    params: list[Any] = []
    conditions: list[str] = []
    param_index = 1

    if search:
        like_conditions = []
        for column in ("article", "manufacturer", "series", "power", "voltage", "ip"):
            like_conditions.append(f"summary.{column} ILIKE ${param_index}")
            params.append(f"%{search}%")
            param_index += 1
        conditions.append("(" + " OR ".join(like_conditions) + ")")

    if filters:
        for column, value in filters.items():
            if column not in available_columns:
                continue
            if value in (None, ""):
                continue
            conditions.append(f"summary.{column} = ${param_index}")
            params.append(value)
            param_index += 1

    where_clause = ""
    if conditions:
        where_clause = " WHERE " + " AND ".join(conditions)

    order_expression = "summary.total_quantity" if sort_column == "total_quantity" else f"summary.{sort_column}"

    query = (
        POWER_SUPPLY_SUMMARY_CTE
        + "\nSELECT * FROM summary"
        + where_clause
        + f" ORDER BY {order_expression} {order_direction.upper()}, summary.article ASC"
    )

    records = await conn.fetch(query, *params)
    rows = [dict(record) for record in records]
    return rows, sort_column, order_direction


async def fetch_power_supply_filter_options(conn: asyncpg.Connection) -> dict[str, list[dict[str, str]]]:
    options: dict[str, list[dict[str, str]]] = {}
    for column in ("manufacturer", "series", "power", "voltage", "ip"):
        records = await conn.fetch(
            POWER_SUPPLY_SUMMARY_CTE
            + f"\nSELECT DISTINCT summary.{column} AS value FROM summary"
            + f" WHERE summary.{column} IS NOT NULL"
            + f" ORDER BY LOWER(summary.{column})"
        )
        options[column] = [
            {"value": record["value"], "label": _format_value(record["value"])}
            for record in records
            if record["value"] is not None
        ]
    return options


@app.get("/materials", response_class=HTMLResponse, name="materials_page")
async def materials_page(
    request: Request,
    search: str = "",
    material: str = "",
    color: str = "",
    thickness: str = "",
    sort: str | None = None,
    order: str | None = None,
    message: str | None = None,
    error: str | None = None,
) -> HTMLResponse:
    filters = {
        "material": material,
        "color": color,
        "thickness": thickness,
    }
    return await _render_table(
        request,
        "materials",
        "materials.html",
        search,
        sort,
        order,
        message,
        error,
        filters=filters,
        filter_columns=["material", "color", "thickness"],
    )


@app.get("/films", response_class=HTMLResponse, name="films_page")
async def films_page(
    request: Request,
    search: str = "",
    sort: str | None = None,
    order: str | None = None,
    message: str | None = None,
    error: str | None = None,
) -> HTMLResponse:
    return await _render_table(request, "films", "films.html", search, sort, order, message, error)


@app.get("/power-supplies", response_class=HTMLResponse, name="power_supplies_page")
async def power_supplies_page(
    request: Request,
    search: str = "",
    manufacturer: str = "",
    series: str = "",
    power: str = "",
    voltage: str = "",
    ip: str = "",
    sort: str | None = None,
    order: str | None = None,
    message: str | None = None,
    error: str | None = None,
) -> HTMLResponse:
    filters = {
        "manufacturer": manufacturer,
        "series": series,
        "power": power,
        "voltage": voltage,
        "ip": ip,
    }

    pool = await get_pool()
    async with pool.acquire() as conn:
        rows, applied_sort, applied_order = await fetch_power_supply_summary(
            conn,
            search,
            sort,
            order,
            filters,
        )
        filter_options = await fetch_power_supply_filter_options(conn)

    export_url = request.url_for("export_power_supplies")
    query_params: dict[str, str] = {}
    if search:
        query_params["search"] = search
    if applied_sort:
        query_params["sort"] = applied_sort
    if applied_order:
        query_params["order"] = applied_order
    for key, value in filters.items():
        if value not in (None, ""):
            query_params[key] = value
    if query_params:
        export_url = f"{export_url}?{urlencode(query_params)}"

    context = {
        "request": request,
        "columns": POWER_SUPPLY_COLUMNS,
        "rows": rows,
        "search": search,
        "sort": applied_sort,
        "order": applied_order,
        "message": message,
        "error": error,
        "export_url": export_url,
        "filters": filters,
        "filter_options": filter_options,
        "route_name": "power_supplies_page",
    }
    return templates.TemplateResponse("power_supplies.html", context)


@app.get("/add", response_class=HTMLResponse, name="add_item_form")
async def add_item_form(
    request: Request,
    type: str = "materials",
    message: str | None = None,
    error: str | None = None,
) -> HTMLResponse:
    item_type = type if type in TABLE_ALIASES else "materials"
    pool = await get_pool()
    columns: list[dict[str, Any]] = []
    try:
        async with pool.acquire() as conn:
            table_info = await resolve_table(conn, item_type)
            raw_columns = await fetch_columns(conn, table_info["schema"], table_info["table"])
            columns = _prepare_form_columns(raw_columns)
    except HTTPException as exc:
        if exc.status_code != status.HTTP_404_NOT_FOUND:
            raise
        error = exc.detail

    context = {
        "request": request,
        "columns": columns,
        "item_type": item_type,
        "message": message,
        "error": error,
    }
    return templates.TemplateResponse("add_item.html", context)


@app.post("/add")
async def add_item(
    request: Request,
    item_type: str = Form(...),
) -> Response:
    item_type = item_type if item_type in TABLE_ALIASES else "materials"
    form = await request.form()
    pool = await get_pool()
    async with pool.acquire() as conn:
        table_info = await resolve_table(conn, item_type)
        raw_columns = await fetch_columns(conn, table_info["schema"], table_info["table"])
        columns_for_form = _prepare_form_columns(raw_columns)

        values: list[Any] = []
        columns_to_insert: list[str] = []
        placeholders: list[str] = []
        param_idx = 1

        for column in raw_columns:
            name = column["name"]
            if name == "id":
                continue
            raw_value = form.get(name)
            try:
                converted = convert_value(raw_value, column["data_type"])
            except (ValueError, TypeError) as exc:
                context = {
                    "request": request,
                    "columns": columns_for_form,
                    "item_type": item_type,
                    "error": f"Ошибка в поле '{column['title']}': {exc}",
                    "message": None,
                }
                return templates.TemplateResponse("add_item.html", context, status_code=400)

            if converted is None:
                if not column["is_nullable"] and not column["has_default"]:
                    context = {
                        "request": request,
                        "columns": columns_for_form,
                        "item_type": item_type,
                        "error": f"Поле '{column['title']}' обязательно для заполнения.",
                        "message": None,
                    }
                    return templates.TemplateResponse("add_item.html", context, status_code=400)
                continue

            columns_to_insert.append(name)
            values.append(converted)
            placeholders.append(f"${param_idx}")
            param_idx += 1

        if not columns_to_insert:
            context = {
                "request": request,
                "columns": columns_for_form,
                "item_type": item_type,
                "error": "Не заполнено ни одно поле для сохранения.",
                "message": None,
            }
            return templates.TemplateResponse("add_item.html", context, status_code=400)

        insert_query = (
            f"INSERT INTO {table_info['qualified']} "
            f"({', '.join(quote_ident(col) for col in columns_to_insert)}) "
            f"VALUES ({', '.join(placeholders)})"
        )
        await conn.execute(insert_query, *values)

    target_route = "materials_page" if item_type == "materials" else "films_page"
    target_url = request.url_for(target_route)
    redirect_url = f"{target_url}?" + urlencode({"message": "Запись успешно добавлена."})
    return RedirectResponse(redirect_url, status_code=status.HTTP_303_SEE_OTHER)


def _create_workbook(columns: list[dict[str, Any]], rows: list[dict[str, Any]]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append([col["title"] for col in columns])

    for row in rows:
        ws.append([row.get(col["name"]) for col in columns])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@app.get("/power-supplies/export", response_class=StreamingResponse, name="export_power_supplies")
async def export_power_supplies(
    search: str = "",
    manufacturer: str = "",
    series: str = "",
    power: str = "",
    voltage: str = "",
    ip: str = "",
    sort: str | None = None,
    order: str | None = None,
) -> StreamingResponse:
    filters = {
        "manufacturer": manufacturer,
        "series": series,
        "power": power,
        "voltage": voltage,
        "ip": ip,
    }

    pool = await get_pool()
    async with pool.acquire() as conn:
        rows, _, _ = await fetch_power_supply_summary(conn, search, sort, order, filters)

    file_buffer = _create_workbook(POWER_SUPPLY_COLUMNS, rows)
    filename = "power-supplies.xlsx"
    headers = {"Content-Disposition": f"attachment; filename={filename}"}
    return StreamingResponse(
        file_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@app.get("/{table_alias}/export", response_class=StreamingResponse, name="export_table")
async def export_table(
    table_alias: str,
    search: str = "",
    sort: str | None = None,
    order: str | None = None,
) -> StreamingResponse:
    if table_alias not in TABLE_ALIASES:
        raise HTTPException(status_code=404, detail="Неизвестный раздел")
    pool = await get_pool()
    async with pool.acquire() as conn:
        table_info = await resolve_table(conn, table_alias)
        columns = await fetch_columns(conn, table_info["schema"], table_info["table"])
        rows, _, _ = await fetch_rows(conn, table_info, columns, search, sort, order)

    file_buffer = _create_workbook(columns, rows)
    filename = f"{table_alias}.xlsx"
    headers = {"Content-Disposition": f"attachment; filename={filename}"}
    return StreamingResponse(file_buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


